"""
Собирает ВСЕ тесты Django-проекта в один Word-файл.

Использование:
    python manage.py export_tests_docx
    python manage.py export_tests_docx --output report.docx
    python manage.py export_tests_docx --apps crud blog
    python manage.py export_tests_docx --no-run     # только код, без запуска
"""
import inspect
import io
import unittest
from datetime import datetime

from django.apps import apps as django_apps
from django.conf import settings
from django.core.management.base import BaseCommand
from django.test.runner import DiscoverRunner
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


def iter_tests(suite):
    for item in suite:
        if isinstance(item, unittest.TestSuite):
            yield from iter_tests(item)
        else:
            yield item


def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


class Command(BaseCommand):
    help = "Экспортирует все тесты проекта (код + результат запуска) в .docx"

    def add_arguments(self, parser):
        parser.add_argument("--output", default="tests_report.docx")
        parser.add_argument("--apps", nargs="*", help="Лейблы приложений (по умолчанию — все локальные)")
        parser.add_argument("--no-run", action="store_true", help="Не запускать тесты, только выгрузить код")
        parser.add_argument("--title", default="Отчёт по тестам Django-проекта")

    # --- поиск приложений проекта (не django.contrib и не сторонние) ---
    def local_labels(self):
        base = str(settings.BASE_DIR)
        labels = []
        for cfg in django_apps.get_app_configs():
            if str(cfg.path).startswith(base) and "site-packages" not in str(cfg.path):
                labels.append(cfg.label)
        return labels

    def handle(self, *args, **opts):
        labels = opts["apps"] or self.local_labels()

        runner = DiscoverRunner(verbosity=0, interactive=False)
        suite = runner.build_suite(labels)
        tests = list(iter_tests(suite))

        results = {}
        if not opts["no_run"] and tests:
            results = self.run_all(runner, suite)

        self.write_docx(opts["output"], opts["title"], tests, results, opts["no_run"])
        self.stdout.write(self.style.SUCCESS(f"Готово: {opts['output']} (тестов: {len(tests)})"))

    # --- запуск с результатами по каждому тесту ---
    def run_all(self, runner, suite):
        runner.setup_test_environment()
        old = runner.setup_databases()
        try:
            res = unittest.TextTestRunner(stream=io.StringIO(), verbosity=0).run(suite)
        finally:
            runner.teardown_databases(old)
            runner.teardown_test_environment()

        out = {}
        for test, tb in res.failures:
            out[test.id()] = ("FAIL", tb)
        for test, tb in res.errors:
            out[test.id()] = ("ERROR", tb)
        for test, _ in res.skipped:
            out[test.id()] = ("SKIP", "")
        return out

    # --- Word ---
    def write_docx(self, path, title, tests, results, no_run):
        doc = Document()
        doc.styles["Normal"].font.name = "Calibri"
        doc.styles["Normal"].font.size = Pt(11)

        h = doc.add_heading(title, 0)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        meta = doc.add_paragraph(f"Дата: {datetime.now():%d.%m.%Y %H:%M}")
        meta.alignment = WD_ALIGN_PARAGRAPH.CENTER

        total = len(tests)
        bad = sum(1 for t in tests if results.get(t.id(), ("OK",))[0] in ("FAIL", "ERROR"))
        if no_run:
            doc.add_paragraph(f"Всего тестов: {total} (запуск не выполнялся)")
        else:
            doc.add_paragraph(f"Всего тестов: {total}. Пройдено: {total - bad}. Упало: {bad}.")

        # сводная таблица
        doc.add_heading("Сводка", 1)
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        for cell, text in zip(table.rows[0].cells, ["№", "Приложение", "Тест", "Результат"]):
            cell.text = text
            shade(cell, "D9E2F3")
            for r in cell.paragraphs[0].runs:
                r.bold = True
        for i, t in enumerate(tests, 1):
            status = "—" if no_run else results.get(t.id(), ("OK",))[0]
            row = table.add_row().cells
            row[0].text = str(i)
            row[1].text = t.__class__.__module__.split(".")[0]
            row[2].text = f"{t.__class__.__name__}.{t._testMethodName}"
            row[3].text = status
            if status == "OK":
                shade(row[3], "E2EFDA")
            elif status in ("FAIL", "ERROR"):
                shade(row[3], "F8CBAD")

        # подробности по каждому тесту
        doc.add_heading("Подробно", 1)
        current_class = None
        for t in tests:
            cls = t.__class__
            if cls is not current_class:
                doc.add_heading(f"{cls.__module__}.{cls.__name__}", 2)
                current_class = cls
            method = getattr(cls, t._testMethodName)
            doc.add_heading(t._testMethodName, 3)

            description = inspect.getdoc(method)
            if description:
                doc.add_paragraph(description)

            try:
                source = inspect.getsource(method)
            except (OSError, TypeError):
                source = "# исходный код недоступен"
            p = doc.add_paragraph()
            run = p.add_run(self.dedent(source))
            run.font.name = "Consolas"
            run.font.size = Pt(9)

            if not no_run:
                status, tb = results.get(t.id(), ("OK", ""))
                sp = doc.add_paragraph()
                r = sp.add_run(f"Результат: {status}")
                r.bold = True
                r.font.color.rgb = RGBColor(0x2E, 0x7D, 0x32) if status == "OK" else RGBColor(0xC6, 0x28, 0x28)
                if tb:
                    tp = doc.add_paragraph()
                    tr = tp.add_run(tb)
                    tr.font.name = "Consolas"
                    tr.font.size = Pt(8)

        doc.save(path)

    @staticmethod
    def dedent(source):
        import textwrap
        return textwrap.dedent(source).rstrip()
